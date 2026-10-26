"""Local web server: JSON API for the UI + serves the built frontend."""

import asyncio
import io
import time
import ipaddress
import os
import re
import socket
import subprocess
import sys
import uuid
import webbrowser

import psutil
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, StrictFloat, StrictInt
from urllib.parse import urlsplit
from .control import Control

from .client import BACnetError, BacnetClient, ConnectError, Device, error_text, probe_port
from .netmon import install_tap, monitor
from .trends import TrendSession
from . import updater
from .version import __version__

app = FastAPI(title="BACon")
client = BacnetClient()
install_tap()
trend = TrendSession(client)
devices: dict[int, Device] = {}
objects: dict[int, dict[str, dict]] = {}  # device instance -> object key -> row
jobs: dict[str, dict] = {}
job_tasks: dict[str, asyncio.Task] = {}
control = Control(client, devices)


@app.middleware('http')
async def protect_control(request: Request, call_next):
    if request.url.path.startswith('/api/control') and request.method == 'POST':
        from fastapi.responses import JSONResponse
        origin = request.headers.get('origin')
        allowed = {request.headers.get('host'), '127.0.0.1:5173', 'localhost:5173'}
        if request.headers.get('x-bacon-control') != '1' or (origin and urlsplit(origin).netloc not in allowed):
            return JSONResponse({'detail': 'Control requests must come from BACon.'}, status_code=403)
    return await call_next(request)


class ControlTarget(BaseModel):
    device: int
    typeId: int
    instance: int
    key: str
    property: str | None = None  # present-value when left out


class ControlPrepare(BaseModel):
    targets: list[ControlTarget] = Field(max_length=2000)


class ControlWrite(BaseModel):
    token: str
    # a number for present value; any JSON the property's kind takes for the rest (see props.py)
    value: StrictFloat | StrictInt | bool | str | list | dict | None = None
    priority: StrictInt | None = None
    release: bool = False


@app.post('/api/control/prepare')
async def control_prepare(body: ControlPrepare):
    return await control.prepare([t.model_dump() for t in body.targets])


@app.post('/api/control/write')
async def control_write(body: ControlWrite):
    try:
        return await control.execute(body.token, body.value, body.priority, body.release)
    except BACnetError as e:
        raise _fail(e)
    except ValueError as e:
        raise HTTPException(409, str(e))


class ControlAck(BaseModel):
    target: ControlTarget
    transition: StrictInt


@app.post('/api/control/acknowledge')
async def control_acknowledge(body: ControlAck):
    try:
        return await control.acknowledge(body.target.model_dump(), body.transition)
    except BACnetError as e:
        raise _fail(e)
    except ValueError as e:
        raise HTTPException(409, str(e))


@app.get('/api/control/actions')
def control_actions(limit: int = 500):
    return control.actions(max(1, min(limit, 5000)))


@app.get('/api/control/writable')
def control_writable():
    """What BACon can write on each object type, {typeId: {property: {kind, enum?, bits?}}}, the enumerations' choices
    once each, and the properties that make up an alarm."""
    from .props import ALARM_PROPS, WRITABLE, describe
    types, enums = {}, {}
    for type_id in range(0, 64):
        found = {}
        for p in sorted(WRITABLE):
            d = describe(type_id, p)
            if d is not None:
                if 'enum' in d:
                    enums[d['enum']] = d.pop('choices')
                found[p] = d
        if found:
            types[type_id] = found
    return {'types': types, 'enums': enums, 'alarm': ALARM_PROPS}


def _fail(err) -> HTTPException:
    if isinstance(err, ConnectError):
        return HTTPException(400, str(err))
    return HTTPException(502, error_text(err))


# ---- network ------------------------------------------------------------


def _not_this_pc() -> set[str]:
    """Addresses on this PC that are not its own: added with SkipAsSource, as Rasher does for each simulated
    controller. Windows never uses them for the PC's traffic, so BACon (the BMS, on the PC's own address) must not
    either."""
    if sys.platform != "win32":
        return set()
    try:
        out = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command",
             "Get-NetIPAddress -AddressFamily IPv4 | Where-Object SkipAsSource | ForEach-Object IPAddress"],
            capture_output=True, text=True, timeout=5, creationflags=subprocess.CREATE_NO_WINDOW,
        ).stdout
    except (OSError, subprocess.SubprocessError):
        return set()
    return {line.strip() for line in out.splitlines() if line.strip()}


@app.get("/api/interfaces")
def interfaces():
    stats = psutil.net_if_stats()
    nics = psutil.net_if_addrs()
    # the check costs PowerShell's start-up, so only when an adapter has more than one address to tell apart
    skip = _not_this_pc() if any(sum(a.family == 2 for a in addrs) > 1 for addrs in nics.values()) else set()
    out = []
    for name, addrs in nics.items():
        for a in addrs:
            if a.family != 2 or a.address.startswith("127.") or a.address.startswith("169.254.") or a.address in skip:
                continue  # IPv4 only, skip loopback, link-local and addresses that aren't this PC's own
            prefix = ipaddress.IPv4Network(f"0.0.0.0/{a.netmask or '255.255.255.0'}").prefixlen
            out.append({
                "name": name,
                "ip": a.address,
                "prefix": prefix,
                "up": bool(stats.get(name) and stats[name].isup),
            })
    return out


def _udp_holders(ip: str, port: int) -> list[dict] | None:
    """Other programs with a UDP socket on this port, on this address or on all addresses.
    None when Windows will not show the socket table, so the UI never claims "nobody else" without knowing."""
    me = os.getpid()
    out, seen = [], set()
    try:
        conns = psutil.net_connections(kind="udp4")
    except (psutil.AccessDenied, OSError):
        return None
    for c in conns:
        if not c.laddr or c.laddr.port != port or c.laddr.ip not in (ip, "0.0.0.0") or c.pid in (None, me, 0):
            continue
        if c.pid in seen:
            continue
        seen.add(c.pid)
        try:
            name = psutil.Process(c.pid).name()
        except (psutil.Error, OSError):
            name = f"process {c.pid}"
        out.append({"name": name, "pid": c.pid, "address": c.laddr.ip})
    return out


def _reserved_range(port: int) -> list[int] | None:
    """The Windows excluded-port range (Hyper-V, WSL, Docker reserve these) that holds this UDP port, if any."""
    if sys.platform != "win32":
        return None
    try:
        out = subprocess.run(
            ["netsh", "int", "ipv4", "show", "excludedportrange", "protocol=udp"],
            capture_output=True, text=True, timeout=3, creationflags=subprocess.CREATE_NO_WINDOW,
        ).stdout
    except (OSError, subprocess.SubprocessError):
        return None
    for lo, hi in re.findall(r"^\s*(\d+)\s+(\d+)", out, re.M):
        if int(lo) <= port <= int(hi):
            return [int(lo), int(hi)]
    return None


@app.get("/api/port-check")
async def port_check(ip: str, port: int = 47808):
    """Lightweight: a test bind like the real one, plus a look at who else holds the port. Nothing is sent.

    holders: other programs on ip:port or 0.0.0.0:port (null when Windows will not say). When BACon is bound
    to ip:port, Windows hands it everything sent to ip:port; a holder on 0.0.0.0 then only gets broadcasts.
    """
    if not 1 <= port <= 65535:
        return {"state": "invalid", "holders": []}
    holders = _udp_holders(ip, port)
    if client.settings and client.settings["ip"] == ip and client.settings["port"] == port:
        return {"state": "ours", "holders": holders}
    stats = psutil.net_if_stats()
    for name, addrs in psutil.net_if_addrs().items():
        if any(a.address == ip for a in addrs) and stats.get(name) and not stats[name].isup:
            return {"state": "down", "holders": []}  # Windows keeps a fixed address on an unplugged adapter
    result = probe_port(ip, port)
    if result == "no-address":
        return {"state": "no-address", "holders": []}
    if result == "in-use":  # only a socket on exactly ip:port blocks the bind; name that one, not the sharers on 0.0.0.0
        return {"state": "busy", "holders": [h for h in holders if h["address"] == ip] or holders if holders else holders}
    if result == "denied":
        # a program holding the port exclusively gives the same error as a reserved range; the socket table tells them apart
        if holders:
            return {"state": "busy", "holders": holders}
        reserved = _reserved_range(port)
        return {"state": "reserved", "range": reserved, "holders": []} if reserved else {"state": "blocked", "holders": holders}
    return {"state": "shared" if holders else "free", "holders": holders}


@app.get("/api/status")
async def status():
    return {"connected": client.app is not None, "settings": client.settings, "host": socket.gethostname()}


class ConnectBody(BaseModel):
    ip: str
    prefix: int = 24
    port: int = 47808
    bbmd: str | None = None


@app.post("/api/connect")
async def connect(body: ConnectBody):
    try:
        client.connect(body.ip, body.prefix, body.port, body.bbmd or None)
    except BACnetError as e:
        raise _fail(e)
    monitor.reset(f"{body.ip}:{body.port}")
    return await status()


@app.post("/api/disconnect")
async def disconnect():
    if trend.running:
        await trend.stop()
    monitor.stop_capture()
    for task in job_tasks.values():
        task.cancel()
    client.close()
    return await status()


# ---- devices --------------------------------------------------------------


def _device_json(d: Device):
    rows = objects.get(d.instance, {})
    data = d.to_json()
    data["objectCount"] = len(rows)
    data["failedCount"] = sum(1 for r in rows.values() if r["status"] == "failed")
    data["partialCount"] = sum(1 for r in rows.values() if r["status"] == "partial")
    return data


class ScanBody(BaseModel):
    low: int | None = None
    high: int | None = None
    timeout: float = 4.0


@app.post("/api/scan")
async def scan(body: ScanBody):
    try:
        found = await client.scan(body.low, body.high, body.timeout)
    except BACnetError as e:
        raise _fail(e)
    for d in found:
        devices[d.instance] = d
    return [_device_json(d) for d in found]


class AddBody(BaseModel):
    address: str


@app.post("/api/devices/add")
async def add_device(body: AddBody):
    try:
        d = await client.add_by_address(body.address)
    except BACnetError as e:
        raise _fail(e)
    devices[d.instance] = d
    return _device_json(d)


@app.get("/api/devices")
async def list_devices():
    return [_device_json(d) for d in sorted(devices.values(), key=lambda d: d.instance)]


@app.delete("/api/devices/{instance}")
async def remove_device(instance: int):
    devices.pop(instance, None)
    objects.pop(instance, None)
    return await list_devices()


@app.get("/api/devices/{instance}/objects")
async def device_objects(instance: int):
    return list(objects.get(instance, {}).values())


@app.get("/api/devices/{instance}/events")
async def device_events(instance: int):
    """What the controller itself reports as in alarm or unacknowledged (GetEventInformation, or GetAlarmSummary)."""
    d = devices.get(instance)
    if d is None:
        raise HTTPException(404, "That device is not in the list any more.")
    if client.app is None:
        raise HTTPException(400, "Connect to a network first.")
    try:
        source, events = await client.events(d)
    except BACnetError as e:
        if "unrecognized-service" in error_text(e):
            name = d.name or f"Device {instance}"
            raise HTTPException(400, f"{name} does not answer alarm queries, so the states shown come from the last read.")
        raise HTTPException(502, _friendly(e))
    return {"source": source, "events": events, "at": time.time()}


# ---- read jobs --------------------------------------------------------------


class ReadBody(BaseModel):
    devices: list[int]
    retryFailed: bool = False
    objects: list[tuple[int, int]] | None = None  # re-read just these (type, instance) pairs


@app.post("/api/read")
async def start_read(body: ReadBody):
    if client.app is None:
        raise HTTPException(400, "Not connected.")
    started = []
    for inst in body.devices:
        d = devices.get(inst)
        if d is None:
            continue
        if any(j["device"] == inst and j["state"] == "running" for j in jobs.values()):
            continue
        only = [tuple(o) for o in body.objects] if body.objects else None
        if body.retryFailed and only is None:
            only = [
                (r["typeId"], r["instance"])
                for r in objects.get(inst, {}).values()
                if r["status"] in ("failed", "partial")
            ]
            if not only:
                continue
        job_id = uuid.uuid4().hex[:8]
        jobs[job_id] = {"id": job_id, "device": inst, "state": "running", "done": 0, "total": 0,
                        "error": "", "retry": body.retryFailed, "only": bool(body.objects)}
        job_tasks[job_id] = asyncio.create_task(_run_read(job_id, d, only))
        started.append(jobs[job_id])
    return started


def _friendly(err) -> str:
    if isinstance(err, ConnectError):
        return str(err)
    text = error_text(err)
    if "no-response" in text or text == "timeout":
        return "The device did not answer. Check it is online and reachable, then try again. Data already read is kept."
    return text


async def _run_read(job_id: str, device: Device, only):
    job = jobs[job_id]
    store = objects.setdefault(device.instance, {})
    # Rows already read stay until fresh ones replace them, so a failed re-read never loses data.
    seen: set[str] = set()

    def progress(done, total, row):
        job["done"], job["total"] = done, total
        if row is not None:
            seen.add(row["key"])
            previous = store.get(row["key"])
            # a failed read never replaces data that was read before
            if not (row["status"] == "failed" and previous and previous["status"] != "failed"):
                store[row["key"]] = row

    try:
        if not device.name and not device.info_error:
            await client.fill_device_info(device)
        await client.read_device(device, progress, only)
        if only is None:
            for key in [k for k in store if k not in seen]:
                del store[key]  # objects the device no longer has
        job["state"] = "done"
    except asyncio.CancelledError:
        job["state"] = "cancelled"
    except BACnetError as e:
        job["state"] = "failed"
        job["error"] = _friendly(e)
    finally:
        job_tasks.pop(job_id, None)


# ---- live values: present value and status flags for the rows on screen ---------------

LIVE_LIMIT = 400
LIVE_PIDS = [85, 111]  # present-value, status-flags
LIVE_FIELDS = {85: "value", 111: "flags"}


class LiveTarget(BaseModel):
    device: int
    typeId: int
    instance: int
    key: str


class LiveBody(BaseModel):
    targets: list[LiveTarget]


@app.post("/api/live")
async def live(body: LiveBody):
    """Read present-value and status-flags again for a few hundred objects. Each value is decoded the way a full read
    decodes it, so it can be compared with the stored row as it is; the stored row keeps up too."""
    if len(body.targets) > LIVE_LIMIT:
        raise HTTPException(400, f"At most {LIVE_LIMIT} objects at a time.")
    if client.app is None:
        raise HTTPException(400, "Connect to a network first.")
    from bacpypes3.primitivedata import ObjectIdentifier
    from bacpypes3.vendor import get_vendor_info

    from .decode import decode_value
    from .names import property_label

    values: dict[str, dict] = {}

    def blank(error):
        return {"value": None, "flags": None, "error": error}

    by_device: dict[int, list[LiveTarget]] = {}
    for t in body.targets:
        by_device.setdefault(t.device, []).append(t)
    reading = {j["device"] for j in jobs.values() if j["state"] == "running"}

    async def batch(inst: int, targets: list[LiveTarget]):
        d = devices.get(inst)
        skip = "unknown device" if d is None else "saved" if d.saved_at else "reading" if inst in reading else None
        if skip:
            for t in targets:
                values[t.key] = blank(skip)
            return
        sem = asyncio.Semaphore(6)
        store = objects.get(inst, {})

        async def one(t: LiveTarget):
            async with sem:
                oid = ObjectIdentifier((t.typeId, t.instance))
                try:
                    # a controller without ReadPropertyMultiple was found out by the full read
                    res = await (client.read_props(d, oid, LIVE_PIDS) if d.no_rpm else client.rpm(d.address, oid, LIVE_PIDS))
                except BACnetError as e:
                    values[t.key] = blank(error_text(e))
                    return
            obj_class = get_vendor_info(0).get_object_class(t.typeId) if t.typeId < 128 else None
            entry = blank(None)
            row = store.get(t.key)
            for pid, field in LIVE_FIELDS.items():
                kind, val = res.get(pid, ("err", "missing"))
                if kind != "ok":
                    if val not in ("unknown-property", "unknownProperty"):  # the object does not have it: not an error
                        entry["error"] = entry["error"] or val
                    continue
                try:
                    entry[field] = decode_value(obj_class, pid, val)
                except Exception as e:
                    entry["error"] = entry["error"] or f"decode failed: {e}"
                    continue
                if row is not None and row["typeId"] == t.typeId and row["instance"] == t.instance:
                    row["values"][property_label(pid, d.vendor_id)] = entry[field]  # the snapshot and the site file keep up
            values[t.key] = entry

        try:
            await asyncio.wait_for(asyncio.gather(*(one(t) for t in targets)), 8.0)
        except asyncio.TimeoutError:
            for t in targets:
                values.setdefault(t.key, blank("timeout"))

    await asyncio.gather(*(batch(inst, targets) for inst, targets in by_device.items()))
    return {"at": time.time(), "values": values}


@app.get("/api/jobs")
async def list_jobs():
    return list(jobs.values())


@app.post("/api/jobs/{job_id}/cancel")
async def cancel_job(job_id: str):
    task = job_tasks.get(job_id)
    if task:
        task.cancel()
    return jobs.get(job_id)


# ---- site files: save everything read, open it again offline --------------------

SITE_FORMAT = "bacon-site"
SITE_FORMATS = (SITE_FORMAT, "bacnet-exporter-site")  # files saved before the rename


@app.get("/api/site/save")
async def save_site(name: str = "site"):
    import datetime
    import gzip
    import json

    saved_at = datetime.datetime.now().isoformat(timespec="seconds")
    data = {
        "format": SITE_FORMAT,
        "version": 1,
        "savedAt": saved_at,
        "host": socket.gethostname(),
        "network": client.settings,
        "devices": [d.to_file() for d in devices.values()],
        "objects": {str(inst): list(rows.values()) for inst, rows in objects.items()},
    }
    body = gzip.compress(json.dumps(data, separators=(",", ":")).encode("utf-8"))
    filename = re.sub(r"[^\w.\- ]", "_", name) or "site"
    return StreamingResponse(
        io.BytesIO(body),
        media_type="application/octet-stream",
        headers={"Content-Disposition": f'attachment; filename="{filename}.bxsite"'},
    )


def _parse_site(raw: bytes) -> dict:
    import gzip
    import json

    try:
        data = json.loads(gzip.decompress(raw) if raw[:2] == bytes([0x1F, 0x8B]) else raw)
    except (OSError, ValueError):
        raise HTTPException(400, "That file is not a BACon site file (.bxsite).")
    if data.get("format") not in SITE_FORMATS:
        raise HTTPException(400, "That file is not a BACon site file (.bxsite).")
    if data.get("version", 0) > 1:
        raise HTTPException(400, "This site file was saved by a newer version of BACon.")
    return data


@app.post("/api/site/open")
async def open_site(request: Request):
    """Merge a saved site into the list. Devices in the file replace the same devices in the list; others stay."""
    data = _parse_site(await request.body())
    saved_at = data.get("savedAt", "")
    loaded = []
    for d in data.get("devices", []):
        device = Device.from_file(d, saved_at)
        if any(j["device"] == device.instance and j["state"] == "running" for j in jobs.values()):
            continue  # never swap data under a read that is still running
        devices[device.instance] = device
        rows = data.get("objects", {}).get(str(device.instance), [])
        objects[device.instance] = {r["key"]: r for r in rows}
        loaded.append(device.instance)
    return {"loaded": loaded, "savedAt": saved_at, "host": data.get("host", "")}


# ---- trends -------------------------------------------------------------------


class TrendAdd(BaseModel):
    objects: list[dict]  # rows: {device, typeId, instance, values}


@app.get("/api/trends")
async def trend_status():
    return trend.status()


@app.post("/api/trends/points")
async def trend_add(body: TrendAdd):
    added = trend.add(body.objects, devices, objects)
    return {"added": added, **trend.status()}


@app.post("/api/trends/points/{point_id}/sync")
async def trend_sync(point_id: int):
    """Read a trend log's records now: after an error, or when it was added while offline."""
    p = trend.points.get(point_id)
    if p is None or p.log is None:
        raise HTTPException(404, "That point is not a trend log.")
    if client.app is None:
        raise HTTPException(400, "Connect to a network first.")
    if p.log.state in ("error", "unsupported", "waiting"):
        p.log.reset()
    await trend.sync_log(p.log)
    return trend.status()


@app.delete("/api/trends/points/{point_id}")
async def trend_remove(point_id: int):
    trend.remove([point_id])
    return trend.status()


class TrendStart(BaseModel):
    interval: float = 10
    mode: str = "auto"


@app.post("/api/trends/start")
async def trend_start(body: TrendStart):
    if client.app is None:
        raise HTTPException(400, "Connect to a network first.")
    if not trend.points:
        raise HTTPException(400, "Add some points to trend first.")
    await trend.start(body.interval, body.mode)
    return trend.status()


@app.post("/api/trends/stop")
async def trend_stop():
    await trend.stop()
    return trend.status()


@app.post("/api/trends/clear")
async def trend_clear():
    if trend.running:
        await trend.stop()
    trend.clear()
    return trend.status()


@app.get("/api/trends/series")
async def trend_series(ids: str, since: float = 0, until: float = 0, buckets: int = 400):
    wanted = [int(x) for x in ids.split(",") if x.strip().isdigit()]
    return trend.series(wanted, since, until, min(max(buckets, 10), 2000))


@app.get("/api/trends/export")
async def trend_export(kind: str = "xlsx"):
    import csv
    import datetime

    rows = trend.export_rows()
    names = {p.id: p for p in trend.points.values()}
    header = ["Time", "Device", "Device name", "Object type", "Instance", "Object name", "Value", "Units"]

    def line(t, pid, v, text):
        p = names.get(pid)
        stamp = datetime.datetime.fromtimestamp(t).isoformat(sep=" ", timespec="seconds")
        if p is None:
            return [stamp, "", "", "", "", f"point {pid}", text, ""]
        from .decode import object_type_name

        return [stamp, str(p.device), p.device_name, object_type_name(p.type_id), str(p.instance), p.name, text, p.units]

    stamp = time.strftime("%Y%m%d-%H%M")
    if kind == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(header)
        for r in rows:
            w.writerow(line(*r))
        data = (chr(0xFEFF) + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly
        return StreamingResponse(io.BytesIO(data), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="bacnet-trend-{stamp}.csv"'})

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Points")
    ws.append(["Device", "Device name", "Object name", "Units", "Method", "Samples", "Min", "Average", "Max", "Last"])
    for p in trend.points.values():
        j = p.to_json()
        ws.append([p.device, p.device_name, p.name, p.units, p.method, p.samples, j["min"], j["avg"], j["max"], p.last])
    ws = wb.create_sheet("Samples")
    ws.append(header)
    for r in rows[:1_000_000]:
        cells = line(*r)
        if r[2] is not None:
            cells[6] = r[2]
        ws.append(cells)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="bacnet-trend-{stamp}.xlsx"'},
    )


# ---- network monitor --------------------------------------------------------------


@app.get("/api/network")
async def network_status(history: int = 300):
    return monitor.status(min(max(history, 60), 3600))


@app.get("/api/network/packets")
async def network_packets(limit: int = 300, address: str = ""):
    return monitor.recent(min(limit, 5000), address)


@app.post("/api/network/capture/start")
async def capture_start():
    monitor.start_capture()
    return monitor.status(60)


@app.post("/api/network/capture/stop")
async def capture_stop():
    monitor.stop_capture()
    return monitor.status(60)


@app.get("/api/network/captures")
async def captures():
    return monitor.captures()


@app.get("/api/network/captures/{name}")
async def capture_download(name: str):
    path = monitor.capture_path(name)
    if not path:
        raise HTTPException(404, "That capture is gone. Captures older than 30 days, or past the newest 10, are removed.")
    if monitor.capture and monitor.capture.path == path:
        monitor.capture.f.flush()
    return FileResponse(path, media_type="application/vnd.tcpdump.pcap", filename=name)


@app.delete("/api/network/captures/{name}")
async def capture_delete(name: str):
    if not monitor.delete_capture(name):
        raise HTTPException(400, "Stop the capture before deleting it.")
    return monitor.captures()


# ---- export ------------------------------------------------------------------


class Sheet(BaseModel):
    name: str
    columns: list[str]
    rows: list[list[str]]


class XlsxBody(BaseModel):
    filename: str = "bacnet-export.xlsx"
    sheets: list[Sheet]


def _cell(value: str):
    # keep numbers as numbers in Excel, everything else as text
    if re.fullmatch(r"-?\d+(\.\d+)?(e[-+]?\d+)?", value or "", re.I) and len(value) < 16:
        try:
            return int(value) if re.fullmatch(r"-?\d+", value) else float(value)
        except ValueError:
            pass
    return value


@app.post("/api/export/xlsx")
def export_xlsx(body: XlsxBody):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for sheet in body.sheets:
        title = re.sub(r"[\[\]:*?/\\]", "_", sheet.name)[:31] or "Sheet"
        base, n = title, 2
        while title in used:
            title = f"{base[:28]}_{n}"
            n += 1
        used.add(title)
        ws = wb.create_sheet(title)
        ws.append(sheet.columns)
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in sheet.rows:
            ws.append([_cell(v) for v in row])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(sheet.columns, 1):
            width = max([len(col)] + [len(r[i - 1]) for r in sheet.rows[:500] if i - 1 < len(r)])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = re.sub(r'[^\w.\- ]', "_", body.filename) or "bacnet-export.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


# ---- version and updates -------------------------------------------------------

# set by the launcher: closes the window and stops the server, so the swap script can replace the exe
on_quit = None


@app.get("/api/version")
def version():
    return {"version": __version__, "app": "BACon", "frozen": updater.FROZEN}


@app.get("/api/update")
async def update_status(check: bool = False, force: bool = False):
    if check or force:
        return await asyncio.to_thread(updater.updater.check, force)
    return updater.updater.status()


@app.post("/api/update/download")
def update_download():
    try:
        updater.updater.download()
    except RuntimeError as e:
        raise HTTPException(400, str(e))
    return updater.updater.status()


@app.post("/api/update/apply")
async def update_apply():
    try:
        updater.updater.apply()
    except RuntimeError as e:
        raise HTTPException(400, str(e))
    if on_quit:
        asyncio.get_running_loop().call_later(0.4, on_quit)
    return {"ok": True}


# ---- frontend ------------------------------------------------------------------


def _web_dir():
    """The built UI: `npm run build` in web/ writes it into the package, so source, pip and the exe all find it here."""
    path = os.path.join(os.path.dirname(__file__), "web")
    return path if os.path.isfile(os.path.join(path, "index.html")) else None


WEB = _web_dir()
if WEB:
    app.mount("/assets", StaticFiles(directory=os.path.join(WEB, "assets")), name="assets")

    @app.get("/{path:path}")
    def spa(path: str):
        target = os.path.join(WEB, path)
        if path and os.path.isfile(target):
            return FileResponse(target)
        return FileResponse(os.path.join(WEB, "index.html"))


def _existing_bacon(port: int) -> bool:
    """Is another BACon already serving this port? Then a second launch just opens a window onto it."""
    import json
    import urllib.request

    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{port}/api/version", timeout=1) as r:
            return json.loads(r.read()).get("app") == "BACon"
    except Exception:
        return False


def _port_free(port: int) -> bool:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        try:
            s.bind(("127.0.0.1", port))
            return True
        except OSError:
            return False


def _log_to_file():
    """The windowed exe has no console, so stdout/stderr are None and uvicorn would crash writing to them."""
    if sys.stdout is None or sys.stderr is None:
        from .paths import data_dir

        log = open(os.path.join(data_dir(), "bacon.log"), "a", encoding="utf-8", buffering=1)
        sys.stdout = sys.stdout or log
        sys.stderr = sys.stderr or log


def main():
    """Start the server, then open BACon in its own window (WebView2). --browser uses the default browser instead;
    --no-browser only runs the server (for development)."""
    global on_quit
    import threading

    import uvicorn

    _log_to_file()
    updater.clean_previous()
    port = int(os.environ.get("BACON_PORT", "8047"))
    server = None
    if _existing_bacon(port):
        pass  # already running: just open another window onto it
    else:
        if not _port_free(port):
            with socket.socket() as s:  # something else has 8047: take any free port
                s.bind(("127.0.0.1", 0))
                port = s.getsockname()[1]
        server = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=port, log_level="warning", log_config=None))
    url = f"http://127.0.0.1:{port}"
    print(f"BACon {__version__} running at {url}")

    mode = "server" if "--no-browser" in sys.argv else "browser" if "--browser" in sys.argv or not WEB else "window"
    if mode == "window":
        try:
            import webview
        except ImportError:
            mode = "browser"

    if mode != "window":
        if mode == "browser" and WEB:
            threading.Timer(1.0, lambda: webbrowser.open(url)).start()
        if server:
            server.run()
        return

    if server:
        threading.Thread(target=server.run, daemon=True).start()
        for _ in range(100):  # wait up to 10 s for the server to answer
            if server.started:
                break
            time.sleep(0.1)

    from .paths import data_dir

    # 1440x900, or 90% of a smaller screen
    screen = webview.screens[0] if webview.screens else None
    width = min(1440, int(screen.width * 0.9)) if screen else 1440
    height = min(900, int(screen.height * 0.9)) if screen else 900
    window = webview.create_window(
        "BACon", url, width=width, height=height, min_size=(960, 600), background_color="#f4f4f5", text_select=True
    )

    def quit_app():
        if server:
            server.should_exit = True
        window.destroy()

    on_quit = quit_app
    webview.settings["ALLOW_DOWNLOADS"] = True
    # a fixed storage folder keeps the layout, saved checks and column choices between runs
    webview.start(private_mode=False, storage_path=data_dir("webview"))
    if server:
        server.should_exit = True


if __name__ == "__main__":
    main()
